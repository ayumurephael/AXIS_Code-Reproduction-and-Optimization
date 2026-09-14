from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "4F79C8"
LIGHT = "D9DEEC"
WHITE = "FFFFFF"
BLACK = "111111"

TABLE_DATA = [
    (
        "GPT-5.5 / Regular",
        {
            "AnomLLM": ("0.2647", "0.3214", "0.8235"),
            "Image LLM": ("0.1176", "0.1538", "0.7647"),
            "LLMAD": ("0.3235", "0.3235", "1.0000"),
        },
    ),
    (
        "DeepSeek / Regular",
        {
            "AnomLLM": ("0.2647", "0.3462", "0.7647"),
            "Image LLM": ("0.0000", "-", "0.0000"),
            "LLMAD": ("0.3235", "0.3235", "1.0000"),
        },
    ),
    (
        "GPT-5.5 / Hard",
        {
            "AnomLLM": ("0.5152", "0.6538", "0.7879"),
            "Image LLM": ("0.4242", "0.5833", "0.7273"),
            "LLMAD": ("0.3636", "0.3750", "0.9697"),
        },
    ),
    (
        "DeepSeek / Hard",
        {
            "AnomLLM": ("0.4545", "0.6522", "0.6970"),
            "Image LLM": ("0.0000", "-", "0.0000"),
            "LLMAD": ("0.4545", "0.4545", "1.0000"),
        },
    ),
]

MODEL_ORDER = ["AnomLLM", "Image LLM", "LLMAD"]


def set_font(run, size: float, *, bold: bool = False, color: str = BLACK) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=150, bottom=90, end=150) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), WHITE)


def set_row_height(row, height_in: float) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_height = OxmlElement("w:trHeight")
    tr_height.set(qn("w:val"), str(int(height_in * 1440)))
    tr_height.set(qn("w:hRule"), "atLeast")
    tr_pr.append(tr_height)


def write_metric_cell(cell, metrics: tuple[str, str, str]) -> None:
    label_acc, parsed_acc, parse_rate = metrics
    lines = [
        f"Label Acc: {label_acc}",
        f"Parsed Acc: {parsed_acc}",
        f"Parse Rate: {parse_rate}",
    ]
    cell.text = ""
    shade_cell(cell, LIGHT)
    set_cell_margins(cell, top=120, start=120, bottom=120, end=120)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for idx, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2 if idx < len(lines) - 1 else 0)
        paragraph.paragraph_format.line_spacing = Pt(15)
        run = paragraph.add_run(line)
        set_font(run, 12.5, color=BLACK)


def build_doc(output_path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.42)
    section.right_margin = Inches(0.42)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    widths = [Inches(2.15), Inches(2.28), Inches(2.28), Inches(2.28)]
    headers = ["评测设置", "AnomLLM", "Image LLM", "LLMAD"]

    header_row = table.rows[0]
    set_row_height(header_row, 0.6)
    for idx, (cell, text) in enumerate(zip(header_row.cells, headers)):
        cell.width = widths[idx]
        shade_cell(cell, BLUE)
        set_cell_margins(cell, top=80, start=80, bottom=80, end=80)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(16)
        r = p.add_run(text)
        set_font(r, 14, bold=True, color=WHITE)

    for label, row_data in TABLE_DATA:
        row = table.add_row()
        set_row_height(row, 1.18)

        label_cell = row.cells[0]
        label_cell.width = widths[0]
        shade_cell(label_cell, BLUE)
        set_cell_margins(label_cell, top=120, start=80, bottom=120, end=80)
        label_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = label_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(16)
        r = p.add_run(label)
        set_font(r, 13.5, bold=True, color=WHITE)

        for col_idx, model_name in enumerate(MODEL_ORDER, start=1):
            row.cells[col_idx].width = widths[col_idx]
            write_metric_cell(row.cells[col_idx], row_data[model_name])

    doc.save(output_path)


def main() -> None:
    output = Path("outputs/reports/baseline_tsdata61_smoke_0606c_metrics_20260610.docx")
    output.parent.mkdir(parents=True, exist_ok=True)
    build_doc(output)
    print(output.resolve())


if __name__ == "__main__":
    main()
