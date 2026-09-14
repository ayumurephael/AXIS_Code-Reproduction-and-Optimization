from __future__ import annotations

import csv
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs" / "local_exports" / "tsdata61_hparam_docx_20260609"
REGULAR_TSV = (
    ROOT
    / "outputs"
    / "studentanswer"
    / "tsdata61_smoke_hparam_eval_20260608a"
    / "regular_smoke_0606c"
    / "metrics_table.tsv"
)
HARD_TSV = (
    ROOT
    / "outputs"
    / "studentanswer"
    / "tsdata61_smoke_hparam_eval_20260608a"
    / "hard_smoke_0606c"
    / "metrics_table.tsv"
)


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def parse_float(text: str) -> float:
    return float(text)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, *, bold: bool = False, size: float = 8.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)


def format_metric(value: str) -> str:
    return f"{parse_float(value):.4f}"


def sort_runs(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    rows = deepcopy(rows)
    rows.sort(
        key=lambda row: (
            -parse_float(row["label_acc"]),
            -parse_float(row["label_acc_on_explicit"]),
            row["run"],
        )
    )
    return rows


def add_dataset_section(
    doc: Document,
    title: str,
    rows: list[dict[str, str]],
    source_path: Path,
) -> None:
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(4)
    run = heading.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(13)

    best_acc_row = max(rows, key=lambda row: parse_float(row["label_acc"]))
    best_parse_row = max(rows, key=lambda row: parse_float(row["label_acc_on_explicit"]))
    summary = doc.add_paragraph()
    summary.paragraph_format.space_after = Pt(6)
    summary.style = doc.styles["Normal"]
    summary.add_run("Best accuracy: ").bold = True
    summary.add_run(
        f"{best_acc_row['run']} ({format_metric(best_acc_row['label_acc'])})"
    )
    summary.add_run("    Best accuracy_on_parse: ").bold = True
    summary.add_run(
        f"{best_parse_row['run']} ({format_metric(best_parse_row['label_acc_on_explicit'])})"
    )

    source = doc.add_paragraph()
    source.paragraph_format.space_after = Pt(6)
    source_run = source.add_run(f"Source: {source_path}")
    source_run.italic = True
    source_run.font.name = "Arial"
    source_run.font.size = Pt(8)
    source_run.font.color.rgb = RGBColor(90, 90, 90)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False

    widths = [Inches(5.2), Inches(1.3), Inches(1.6), Inches(1.3)]
    headers = ["run", "accuracy", "accuracy_on_parse", "parse_rate"]
    hdr_cells = table.rows[0].cells
    for idx, text in enumerate(headers):
        hdr_cells[idx].width = widths[idx]
        set_cell_text(hdr_cells[idx], text, bold=True, size=9)
        shade_cell(hdr_cells[idx], "D9EAF7")
    set_repeat_table_header(table.rows[0])

    best_acc = parse_float(best_acc_row["label_acc"])
    best_parse = parse_float(best_parse_row["label_acc_on_explicit"])

    for row in sort_runs(rows):
        tr = table.add_row().cells
        values = [
            row["run"],
            format_metric(row["label_acc"]),
            format_metric(row["label_acc_on_explicit"]),
            format_metric(row["explicit_answer_rate"]),
        ]
        highlight = (
            parse_float(row["label_acc"]) == best_acc
            or parse_float(row["label_acc_on_explicit"]) == best_parse
        )
        for idx, value in enumerate(values):
            tr[idx].width = widths[idx]
            set_cell_text(tr[idx], value, bold=highlight)
            if highlight:
                shade_cell(tr[idx], "EEF6EA")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(3)


def build_docx(out_path: Path) -> None:
    regular_rows = read_tsv(REGULAR_TSV)
    hard_rows = read_tsv(HARD_TSV)

    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("TSData61 Smoke 0606c Hyperparameter Search Summary")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    subrun = subtitle.add_run(
        "Question sets: question_tsdata61_regular_smoke_0606c / "
        "question_tsdata61_hard_smoke_0606c"
    )
    subrun.font.name = "Arial"
    subrun.font.size = Pt(10)

    note = doc.add_paragraph()
    note.paragraph_format.space_after = Pt(10)
    note_run = note.add_run(
        "Metrics shown: accuracy (label_acc), accuracy_on_parse "
        "(label_acc_on_explicit), and parse_rate (explicit_answer_rate). "
        "The source artifacts are the local synchronized results under "
        "outputs/studentanswer/tsdata61_smoke_hparam_eval_20260608a."
    )
    note_run.font.name = "Arial"
    note_run.font.size = Pt(9)

    generated = doc.add_paragraph()
    generated.paragraph_format.space_after = Pt(8)
    gen_run = generated.add_run(
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )
    gen_run.font.name = "Arial"
    gen_run.font.size = Pt(8)
    gen_run.font.color.rgb = RGBColor(90, 90, 90)

    add_dataset_section(doc, "Regular Smoke 0606c", regular_rows, REGULAR_TSV)
    doc.add_page_break()
    add_dataset_section(doc, "Hard Smoke 0606c", hard_rows, HARD_TSV)

    doc.save(out_path)


def write_combined_tsv(out_path: Path) -> None:
    fields = [
        "split",
        "run",
        "accuracy",
        "accuracy_on_parse",
        "parse_rate",
    ]
    regular_rows = read_tsv(REGULAR_TSV)
    hard_rows = read_tsv(HARD_TSV)

    with out_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, delimiter="\t")
        writer.writeheader()
        for split, rows in (("regular", regular_rows), ("hard", hard_rows)):
            for row in rows:
                writer.writerow(
                    {
                        "split": split,
                        "run": row["run"],
                        "accuracy": row["label_acc"],
                        "accuracy_on_parse": row["label_acc_on_explicit"],
                        "parse_rate": row["explicit_answer_rate"],
                    }
                )


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    combined_tsv = OUT_DIR / "tsdata61_smoke_0606c_hparam_summary.tsv"
    docx_path = OUT_DIR / "tsdata61_smoke_0606c_hparam_summary.docx"
    write_combined_tsv(combined_tsv)
    build_docx(docx_path)
    print(combined_tsv)
    print(docx_path)


if __name__ == "__main__":
    main()